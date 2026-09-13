from __future__ import annotations

import math
from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook
from pydantic import Field, model_validator

from .marlin_contracts import (
    MarlinClassificationDecision,
    MarlinFeatureSummary,
    MarlinGroupedScore,
    MarlinPredictionReport,
    MarlinSourceKind,
)
from .marlin_runtime import MarlinRuntimeResult
from .models import FileFingerprint, GenomeBuild, ModuleRunStatus, StrictModel

_REQUIRED_COLUMNS = ("model_id", "class_name_current", "mcf", "lineage")
_EXPECTED_MODEL_UNITS = 42
_CONFIDENCE_THRESHOLD = 0.8


class MarlinClassAnnotation(StrictModel):
    model_id: int = Field(ge=1, le=42)
    class_name_current: str = Field(min_length=1)
    methylation_class_family: str = Field(min_length=1)
    lineage: str = Field(min_length=1)

    @model_validator(mode="after")
    def text_fields_are_real_labels(self) -> MarlinClassAnnotation:
        for value in (
            self.class_name_current,
            self.methylation_class_family,
            self.lineage,
        ):
            if value.casefold() in {"na", "n/a", "none", "null"}:
                raise ValueError("MARLIN annotation labels cannot be missing-value tokens")
        return self


def _text_cell(value: object, *, column: str, row_number: int) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"MARLIN annotation row {row_number} requires non-empty {column}")
    if value != value.strip():
        raise ValueError(f"MARLIN annotation row {row_number} has padded {column}")
    return value


def load_marlin_class_annotations(path: Path) -> tuple[MarlinClassAnnotation, ...]:
    """Load the locked upstream class annotation workbook by stable ``model_id``."""
    candidate = Path(path)
    if not candidate.is_file():
        raise ValueError(f"MARLIN class annotation workbook is missing: {candidate}")
    workbook = load_workbook(candidate, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration as exc:
            raise ValueError("MARLIN class annotation workbook is empty") from exc
        header = [str(value).strip() if value is not None else "" for value in header_values]
        if len(header) != len(set(header)):
            raise ValueError("MARLIN class annotation workbook contains duplicate column names")
        missing = [column for column in _REQUIRED_COLUMNS if column not in header]
        if missing:
            raise ValueError(
                "MARLIN class annotation workbook is missing required column(s): "
                + ", ".join(missing)
            )
        index = {column: header.index(column) for column in _REQUIRED_COLUMNS}

        annotations: list[MarlinClassAnnotation] = []
        for row_number, row in enumerate(rows, 2):
            if all(value is None for value in row):
                continue
            model_id_value = row[index["model_id"]]
            if isinstance(model_id_value, bool) or not isinstance(model_id_value, int):
                raise ValueError(
                    f"MARLIN annotation row {row_number} model_id must be an integer"
                )
            annotations.append(
                MarlinClassAnnotation(
                    model_id=model_id_value,
                    class_name_current=_text_cell(
                        row[index["class_name_current"]],
                        column="class_name_current",
                        row_number=row_number,
                    ),
                    methylation_class_family=_text_cell(
                        row[index["mcf"]], column="mcf", row_number=row_number
                    ),
                    lineage=_text_cell(
                        row[index["lineage"]], column="lineage", row_number=row_number
                    ),
                )
            )
    finally:
        workbook.close()

    ordered = tuple(sorted(annotations, key=lambda item: item.model_id))
    ids = [item.model_id for item in ordered]
    if ids != list(range(1, _EXPECTED_MODEL_UNITS + 1)):
        raise ValueError("MARLIN class annotations require model IDs 1..42 exactly once")
    return ordered


def _validate_annotations(
    annotations: Sequence[MarlinClassAnnotation],
) -> dict[int, MarlinClassAnnotation]:
    rows = tuple(sorted(annotations, key=lambda item: item.model_id))
    ids = [item.model_id for item in rows]
    if ids != list(range(1, _EXPECTED_MODEL_UNITS + 1)):
        raise ValueError("MARLIN class annotations require model IDs 1..42 exactly once")
    return {item.model_id: item for item in rows}


def _aggregate_scores(
    runtime_result: MarlinRuntimeResult,
    annotation_by_id: dict[int, MarlinClassAnnotation],
    *,
    field: str,
) -> list[MarlinGroupedScore]:
    labels: list[str] = []
    members: dict[str, list[float]] = {}
    for unit in runtime_result.model_scores:
        annotation = annotation_by_id[unit.model_id]
        label = getattr(annotation, field)
        if label not in members:
            labels.append(label)
            members[label] = []
        members[label].append(unit.score)
    return [
        MarlinGroupedScore(label=label, score=math.fsum(members[label])) for label in labels
    ]


def classify_marlin_scores(
    *,
    sample_id: str,
    runtime_result: MarlinRuntimeResult | None,
    feature_summary: MarlinFeatureSummary,
    input_fingerprint: FileFingerprint,
    source_kind: MarlinSourceKind,
    genome_build: GenomeBuild,
    artifact_lock_id: str,
    annotations: Sequence[MarlinClassAnnotation],
) -> MarlinPredictionReport:
    """Bind raw MARLIN units to the locked taxonomy and apply published confidence semantics."""
    if genome_build is not GenomeBuild.GRCH37:
        raise ValueError("MARLIN v1 classification currently requires GRCh37/hg19")
    annotation_by_id = _validate_annotations(annotations)

    if feature_summary.observed_model_feature_count == 0:
        if runtime_result is not None:
            raise ValueError(
                "MARLIN zero-evidence NO_CALL must not execute or carry a runtime result"
            )
        return MarlinPredictionReport(
            sample_id=sample_id,
            status=ModuleRunStatus.NO_CALL,
            decision=MarlinClassificationDecision.UNKNOWN,
            source_kind=source_kind,
            genome_build=genome_build,
            input_fingerprint=input_fingerprint,
            feature_summary=feature_summary,
            raw_model_scores=[],
            class_scores=[],
            family_scores=[],
            lineage_scores=[],
            top_class=None,
            top_class_score=None,
            artifact_lock_id=artifact_lock_id,
            limitations=["No MARLIN model feature was observed; inference was not run."],
        )

    if runtime_result is None:
        raise ValueError("MARLIN observed features require a runtime result")
    if runtime_result.feature_vector_sha256 != feature_summary.feature_vector_sha256:
        raise ValueError(
            "MARLIN runtime feature vector differs from classification feature summary"
        )
    if runtime_result.artifact_lock_id != artifact_lock_id:
        raise ValueError("MARLIN runtime artifact lock differs from classification artifact lock")

    class_scores = _aggregate_scores(
        runtime_result, annotation_by_id, field="class_name_current"
    )
    family_scores = _aggregate_scores(
        runtime_result, annotation_by_id, field="methylation_class_family"
    )
    lineage_scores = _aggregate_scores(runtime_result, annotation_by_id, field="lineage")
    top = max(class_scores, key=lambda item: item.score)
    decision = (
        MarlinClassificationDecision.HIGH_CONFIDENCE
        if top.score >= _CONFIDENCE_THRESHOLD
        else MarlinClassificationDecision.UNKNOWN
    )
    return MarlinPredictionReport(
        sample_id=sample_id,
        status=ModuleRunStatus.COMPLETED,
        decision=decision,
        source_kind=source_kind,
        genome_build=genome_build,
        input_fingerprint=input_fingerprint,
        feature_summary=feature_summary,
        raw_model_scores=runtime_result.model_scores,
        class_scores=class_scores,
        family_scores=family_scores,
        lineage_scores=lineage_scores,
        top_class=top.label,
        top_class_score=top.score,
        artifact_lock_id=artifact_lock_id,
        limitations=[
            "MARLIN classification is Research Use Only and is not a clinical diagnosis."
        ],
    )
