"""The QDNAseq/ACE copy-number lane as a first-class member of the stage graph.

The lane contributes three things to a run and replaces nothing:

* the ``cnv`` stage itself (:func:`plan`, :func:`execute`), configured by the run's
  :class:`CnvLaneSettings` rather than by process-global registration;
* the copy-number part of result assembly (:func:`merge_into_result`): CNV events, the
  recomputed ISCN proposal and the CNV module outcome;
* the copy-number part of the reviewer report (plots through ``render_html`` and the
  workbook sheets added by :func:`enrich_workbook`).

Assembly and reporting themselves stay single implementations in
:mod:`ontseq_platform.pipeline.runner`. Copy-number evidence enters them only through
:func:`load_current_cnv`, i.e. as an artifact this run's CNV stage recorded and that still
verifies byte for byte. A report left in the envelope by an earlier attempt, a failed
re-run or a run that did not request CNV is never merged into a result.
"""

from __future__ import annotations

import json
import sqlite3
from collections.abc import Sequence
from contextlib import closing
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import load_workbook

from ..annotation_cache import require_annotation_cache_build, validate_annotation_cache
from ..iscn import (
    build_iscn_proposal,
    iscn_module_outcome,
    resource_provenance_for_iscn,
)
from ..models import (
    AnalysisModule,
    CraminoQCReport,
    EventType,
    ISCNAssessmentBlocker,
    ISCNResourceProvenance,
    ISCNSelectionPolicy,
    ModuleOutcome,
    ModuleRunStatus,
    PipelineResult,
    Provenance,
    SidecarArtifact,
    Verdict,
)
from ..pipeline.context import StagePlan, StageResult, current_artifact
from ..pipeline.envelope import Artifact, sha256_file
from ..pipeline.stages import StageId
from ..sidecars import tabular_sidecar
from .cytoband import AffectedBandGroup, CnvDirection, CnvSegment, Cytoband, annotate_cnv_cytobands
from .qdnaseq import (
    WHOLE_CHROMOSOME_CONFIRMATION,
    QDNAseqCallReport,
    QDNAseqPolicy,
    run_qdnaseq_ace,
)

if TYPE_CHECKING:
    from ..pipeline.context import RunContext

CNV_DIR = "evidence/cnv/qdnaseq"
CNV_REPORT = "evidence/cnv/{sample}.qdnaseq.json"
CNV_CYTOBAND_REPORT = "evidence/cnv/{sample}.cytobands.json"
#: Identifies this lane in assembly and report plans, and therefore in their resume signatures.
LANE_ID = "qdnaseq-ace-v1"


@dataclass(frozen=True)
class CnvLaneSettings:
    """How this run calls copy number: the versioned policy and the pinned R runner."""

    policy: QDNAseqPolicy
    rscript: str = "Rscript"
    script: Path = Path("scripts/run_qdnaseq_ace.R")


def _requested(ctx: RunContext) -> bool:
    return AnalysisModule.CNV in set(ctx.manifest.analysis.modules)


def _probe_r_packages(ctx: RunContext, settings: CnvLaneSettings) -> dict[str, str]:
    r_version = ctx.runner.run([settings.rscript, "--version"], timeout_seconds=60)
    if r_version.returncode != 0:
        raise ValueError(f"Rscript version probe failed: {r_version.stderr.strip()}")
    expression = (
        "cat(as.character(packageVersion('QDNAseq')), '\\n');"
        "cat(as.character(packageVersion('ACE')), '\\n');"
        "cat(as.character(packageVersion('DNAcopy')), '\\n')"
    )
    packages = ctx.runner.run([settings.rscript, "-e", expression], timeout_seconds=60)
    if packages.returncode != 0:
        diagnostic = packages.stderr.strip()[-2000:]
        raise ValueError(f"QDNAseq/ACE R packages are unavailable: {diagnostic}")
    lines = [line.strip() for line in packages.stdout.splitlines() if line.strip()]
    if len(lines) < 3:
        raise ValueError("could not determine QDNAseq, ACE and DNAcopy versions")
    r_text = r_version.stderr.strip() or r_version.stdout.strip()
    return {
        "Rscript": r_text,
        "QDNAseq": lines[0],
        "ACE": lines[1],
        "DNAcopy": lines[2],
    }


def plan(ctx: RunContext, settings: CnvLaneSettings) -> StagePlan:
    """Plan the QDNAseq/ACE run, or record that the manifest did not ask for copy number."""
    if not _requested(ctx):
        return StagePlan(parameters={"requested": False}, tool_versions={})
    if not settings.script.is_file():
        raise ValueError(f"QDNAseq R runner not found: {settings.script}")
    versions = _probe_r_packages(ctx, settings)
    bam = Path(ctx.manifest.input.path)
    external_inputs = [
        ctx.fingerprint_external_input(bam),
        ctx.fingerprint_external_input(settings.script),
    ]
    if ctx.config.annotation_cache is not None:
        external_inputs.append(
            ctx.fingerprint_external_input(ctx.config.annotation_cache, label="annotation_cache")
        )
    return StagePlan(
        parameters={
            "requested": True,
            "profile": settings.policy.profile_id,
            "bin_sizes_kbp": settings.policy.bin_sizes_kbp,
            "primary_bin_size_kbp": settings.policy.primary_bin_size_kbp,
            "ace_penalty": settings.policy.ace_penalty,
            "ploidy_min": settings.policy.ploidy_min,
            "ploidy_max": settings.policy.ploidy_max,
            "ploidy_step": settings.policy.ploidy_step,
            "cytoband_policy_schema": settings.policy.schema_version,
            "cytoband_policy_id": settings.policy.profile_id,
            "cytoband_affected_fraction": settings.policy.cytoband_affected_fraction,
            "whole_chromosome_fraction": settings.policy.whole_chromosome_fraction,
            "whole_chromosome_span_basis": settings.policy.whole_chromosome_span_basis,
            "whole_chromosome_confirmation": WHOLE_CHROMOSOME_CONFIRMATION,
            "threads": ctx.config.threads,
        },
        tool_versions=versions,
        external_inputs=tuple(external_inputs),
    )


def _load_cytobands(annotation_cache: Path, *, expected_build: str = "GRCh38") -> list[Cytoband]:
    with closing(
        sqlite3.connect(f"file:{annotation_cache.as_posix()}?mode=ro", uri=True)
    ) as connection:
        require_annotation_cache_build(connection, expected_build=expected_build)
        rows = connection.execute(
            "SELECT chrom, start, end, name, gie_stain FROM cytobands "
            "ORDER BY chrom, start, end, name"
        ).fetchall()
    return [
        Cytoband(
            chromosome=str(chromosome),
            start=int(start),
            end=int(end),
            name=str(name),
            gie_stain=str(stain) if stain is not None else None,
        )
        for chromosome, start, end, name, stain in rows
    ]


def _annotate_cnv_cytobands(
    ctx: RunContext,
    report: QDNAseqCallReport,
    settings: CnvLaneSettings,
) -> tuple[QDNAseqCallReport, Artifact | None]:
    cache = ctx.config.annotation_cache
    if cache is None or not report.events:
        return report, None
    directions: dict[EventType, CnvDirection] = {
        EventType.CHROMOSOME_GAIN: "gain",
        EventType.DUPLICATION: "gain",
        EventType.CHROMOSOME_LOSS: "loss",
        EventType.DELETION: "loss",
    }
    segments = [
        CnvSegment(
            event_id=event.event_id,
            chromosome=event.primary.chromosome,
            start=event.primary.start,
            end=event.primary.end,
            direction=directions[event.event_type],
            whole_chromosome=event.event_type
            in {EventType.CHROMOSOME_GAIN, EventType.CHROMOSOME_LOSS},
        )
        for event in report.events
        if event.event_type in directions
    ]
    chromosome_sizes = {item.name: item.length for item in ctx.config.reference_lock.contigs}
    annotation = annotate_cnv_cytobands(
        segments,
        _load_cytobands(cache, expected_build=ctx.config.manifest.assay.genome_build.value),
        affected_fraction=settings.policy.cytoband_affected_fraction,
        chromosome_sizes=chromosome_sizes,
    )
    affected_groups: dict[str, list[AffectedBandGroup]] = {}
    for group in annotation.affected_groups:
        for event_id in group.source_event_ids:
            affected_groups.setdefault(event_id, []).append(group)
    events = []
    for event in report.events:
        unannotated_locus = event.primary.model_copy(
            update={"cytoband_start": None, "cytoband_end": None}
        )
        groups = sorted(
            affected_groups.get(event.event_id, []),
            key=lambda value: (value.start, value.end, value.start_band, value.end_band),
        )
        if not groups:
            events.append(event.model_copy(update={"primary": unannotated_locus}))
            continue
        if len(groups) != 1 or groups[0].source_event_ids != (event.event_id,):
            source_count = len(
                set(event_id for group in groups for event_id in group.source_event_ids)
            )
            events.append(
                event.model_copy(
                    update={
                        "primary": unannotated_locus,
                        "notes": [
                            *event.notes,
                            f"Cytoband projection produced {len(groups)} group(s) spanning "
                            f"{source_count} source event(s); no single-event ISCN band range "
                            "was assigned.",
                        ],
                    }
                )
            )
            continue
        start_band = groups[0].start_band
        end_band = groups[0].end_band
        locus = unannotated_locus.model_copy(
            update={"cytoband_start": start_band, "cytoband_end": end_band}
        )
        events.append(
            event.model_copy(
                update={
                    "primary": locus,
                    "notes": [
                        *event.notes,
                        "Cytoband affected-fraction threshold "
                        f"{annotation.threshold:g}: {start_band}-{end_band}.",
                    ],
                }
            )
        )
    payload = {
        "schema_version": "1.0.0",
        "genome_build": ctx.config.manifest.assay.genome_build.value,
        "policy": {
            "schema_version": settings.policy.schema_version,
            "profile_id": settings.policy.profile_id,
            "cytoband_affected_fraction": settings.policy.cytoband_affected_fraction,
        },
        "threshold": annotation.threshold,
        "raw_overlaps": [asdict(item) for item in annotation.raw_overlaps],
        "affected_groups": [asdict(item) for item in annotation.affected_groups],
        "whole_chromosome_calls": [asdict(item) for item in annotation.whole_chromosome_calls],
    }
    artifact = ctx.envelope.atomic_write_text(
        ctx.path(CNV_CYTOBAND_REPORT), json.dumps(payload, indent=2) + "\n"
    )
    annotated_report = QDNAseqCallReport.model_validate(
        report.model_copy(update={"events": events}).model_dump(mode="python")
    )
    return annotated_report, artifact


def _record_cnv_outputs(ctx: RunContext, report: QDNAseqCallReport) -> list[Artifact]:
    artifacts: list[Artifact] = []
    relative_report = ctx.path(CNV_REPORT)
    artifacts.append(
        ctx.envelope.atomic_write_text(relative_report, report.model_dump_json(indent=2) + "\n")
    )
    for name in report.output_files:
        if name.lower().endswith(".rds"):
            continue
        artifacts.append(ctx.envelope.fingerprint(f"{CNV_DIR}/{name}"))
    return artifacts


def execute(ctx: RunContext, stage_plan: StagePlan, settings: CnvLaneSettings) -> StageResult:
    """Run QDNAseq/ACE and record every normalized output as a checksummed stage artifact."""
    del stage_plan
    if not _requested(ctx):
        return StageResult(
            status=ModuleRunStatus.NOT_RUN,
            reason="CNV was not requested in the sample manifest.",
        )
    # A re-execution owns the normalized report path. Remove it first, so a failed attempt
    # cannot leave an earlier attempt's report looking like this run's output on disk.
    for template in (CNV_REPORT, CNV_CYTOBAND_REPORT):
        ctx.envelope.path(ctx.path(template)).unlink(missing_ok=True)
    report = run_qdnaseq_ace(
        bam=Path(ctx.manifest.input.path),
        sample_id=ctx.sample_id,
        genome_build=ctx.manifest.assay.genome_build,
        reference_lock=ctx.config.reference_lock,
        policy=settings.policy,
        output_dir=ctx.envelope.path(CNV_DIR),
        script=settings.script,
        runner=ctx.runner,
        rscript=settings.rscript,
        threads=ctx.config.threads,
    )
    report, cytoband_artifact = _annotate_cnv_cytobands(ctx, report, settings)
    artifacts = _record_cnv_outputs(ctx, report)
    if cytoband_artifact is not None:
        artifacts.append(cytoband_artifact)
    bins = ", ".join(str(value) for value in settings.policy.bin_sizes_kbp)
    reason = (
        f"QDNAseq+ACE completed at {bins} kbp; primary "
        f"{report.primary_fit.bin_size_kbp} kbp fit: cellularity "
        f"{report.primary_fit.cellularity:.3f}, ploidy {report.primary_fit.ploidy:.3f}; "
        f"{len(report.events)} normalized CNV event(s)."
    )
    return StageResult(
        status=report.status,
        reason=reason,
        outputs=artifacts,
        tools=report.tools,
        warnings=report.warnings,
        limitations=report.limitations,
    )


def load_current_cnv(ctx: RunContext) -> QDNAseqCallReport | None:
    """Read the copy-number report this run's CNV stage recorded, if it recorded one.

    The only way copy-number evidence reaches assembly or the reviewer report. A report
    that merely exists in the envelope, from an earlier attempt, a failed re-run or a run
    that never requested CNV, is not evidence about this run.
    """
    relative_path = ctx.path(CNV_REPORT)
    if current_artifact(ctx, StageId.CNV, relative_path) is None:
        return None
    return QDNAseqCallReport.model_validate_json(
        ctx.envelope.path(relative_path).read_text(encoding="utf-8")
    )


def _verified_iscn_resource_provenance(
    ctx: RunContext,
) -> tuple[ISCNResourceProvenance | None, list[ISCNAssessmentBlocker]]:
    context = ctx.config.resource_context
    provenance = resource_provenance_for_iscn(context)
    cache = ctx.config.annotation_cache
    if provenance is None or context is None or cache is None:
        return None, [
            ISCNAssessmentBlocker(
                reason_code="RESOURCE_PROVENANCE_UNAVAILABLE",
                detail="resolved reference/cytoband provenance or annotation cache is unavailable",
            )
        ]
    try:
        configured_cache = cache.resolve(strict=True)
        context_cache = Path(context.resource_paths["reference.annotation_cache"]).resolve(
            strict=True
        )
        reference_lock = Path(context.resource_paths["reference.reference_lock"])
        cytobands = Path(context.resource_paths["reference.cytobands"])
        cache_summary = validate_annotation_cache(
            configured_cache,
            expected_build=ctx.manifest.assay.genome_build.value,
        )
        with closing(
            sqlite3.connect(f"file:{configured_cache.as_posix()}?mode=ro", uri=True)
        ) as connection:
            require_annotation_cache_build(
                connection,
                expected_build=ctx.manifest.assay.genome_build.value,
            )
            metadata = dict(connection.execute("SELECT key, value FROM metadata").fetchall())
        reference_lock_sha256 = sha256_file(reference_lock)
        cytoband_sha256 = sha256_file(cytobands)
    except (KeyError, OSError, sqlite3.DatabaseError, ValueError) as exc:
        return None, [
            ISCNAssessmentBlocker(
                reason_code="RESOURCE_PROVENANCE_UNREADABLE",
                detail=f"ISCN reference/cytoband/cache provenance could not be verified: {exc}",
            )
        ]
    expected = {
        "bundle_id": provenance.reference_bundle_id,
        "bundle_version": provenance.reference_bundle_version,
        "cytoband_release": provenance.cytoband_release,
        "cytoband_sha256": provenance.cytoband_sha256,
    }
    mismatches = []
    if configured_cache != context_cache:
        mismatches.append(
            "configured annotation cache is not the cache selected by the resolved resource context"
        )
    actual_resource_sha256 = {
        "reference_lock_sha256": reference_lock_sha256,
        "cytoband_sha256": cytoband_sha256,
        "annotation_cache_sha256": cache_summary.sha256,
    }
    expected_resource_sha256 = {
        "reference_lock_sha256": provenance.reference_lock_sha256,
        "cytoband_sha256": provenance.cytoband_sha256,
        "annotation_cache_sha256": provenance.annotation_cache_sha256,
    }
    mismatches.extend(
        f"{name}: expected {expected_sha!r}, observed {actual_resource_sha256[name]!r}"
        for name, expected_sha in expected_resource_sha256.items()
        if actual_resource_sha256[name] != expected_sha
    )
    mismatches.extend(
        f"{key}: expected {value!r}, observed {metadata.get(key)!r}"
        for key, value in expected.items()
        if metadata.get(key) != value
    )
    if mismatches:
        return None, [
            ISCNAssessmentBlocker(
                reason_code="RESOURCE_CACHE_PROVENANCE_MISMATCH",
                detail="; ".join(mismatches),
            )
        ]
    return provenance, []


def assemble_parameters(settings: CnvLaneSettings) -> dict[str, object]:
    """What this lane adds to the assembly plan, and therefore to its resume signature."""
    return {
        "cnv_lane": LANE_ID,
        "iscn_whole_chromosome_fraction": settings.policy.whole_chromosome_fraction,
        "iscn_cytoband_affected_fraction": settings.policy.cytoband_affected_fraction,
    }


def _replace_module(
    modules: Sequence[ModuleOutcome],
    replacement: ModuleOutcome,
) -> list[ModuleOutcome]:
    result = [item for item in modules if item.module != replacement.module]
    result.append(replacement)
    return sorted(result, key=lambda item: item.module.value)


def _merge_provenance(base: Provenance, cnv: QDNAseqCallReport) -> Provenance:
    return base.model_copy(update={"tools": [*base.tools, *cnv.tools]})


def merge_into_result(
    ctx: RunContext,
    result: PipelineResult,
    *,
    qc: CraminoQCReport,
    cnv: QDNAseqCallReport,
    settings: CnvLaneSettings,
) -> PipelineResult:
    """Add current copy-number evidence to an assembled result and recompute ISCN."""
    if cnv.sample_id != ctx.manifest.sample_id:
        raise ValueError("Manifest and QDNAseq/ACE artifact must refer to the same sample")
    if cnv.genome_build != ctx.manifest.assay.genome_build:
        raise ValueError("Manifest and QDNAseq/ACE artifact use different genome builds")
    sidecars: list[SidecarArtifact] = list(result.sidecars)
    primary_tables = (
        ("cnv_bins", cnv.primary_fit.bins_file),
        ("cnv_segments", cnv.primary_fit.segment_file),
        ("ace_models", cnv.primary_fit.model_file),
    )
    for artifact_id, filename in primary_tables:
        if filename is None:
            continue
        sidecars.append(
            tabular_sidecar(
                artifact_id=artifact_id,
                envelope_root=ctx.envelope.root,
                relative_path=f"{CNV_DIR}/{filename}",
            )
        )
    outcome = ModuleOutcome(
        module=AnalysisModule.CNV,
        status=cnv.status,
        reason=(
            f"QDNAseq+ACE multi-bin CNV completed; primary "
            f"{cnv.primary_fit.bin_size_kbp} kbp, cellularity "
            f"{cnv.primary_fit.cellularity:.3f}, ploidy {cnv.primary_fit.ploidy:.3f}."
        ),
        tools=cnv.tools,
    )
    policy = settings.policy
    combined_events = [*cnv.events, *result.events]
    iscn_provenance, iscn_blockers = _verified_iscn_resource_provenance(ctx)
    if qc.qc.verdict == Verdict.FAIL:
        iscn_blockers.append(
            ISCNAssessmentBlocker(
                reason_code="QC_FAILED",
                detail=(
                    "run QC failed; candidate events remain visible but ISCN assessment is blocked"
                ),
            )
        )
    proposal = build_iscn_proposal(
        combined_events,
        genome_build=ctx.manifest.assay.genome_build,
        selection_policy=ISCNSelectionPolicy.TECHNICAL_CANDIDATES_V1,
        requested=AnalysisModule.ISCN in set(ctx.manifest.analysis.modules),
        upstream_evidence_assessed=cnv.status
        in {ModuleRunStatus.COMPLETED, ModuleRunStatus.NO_CALL},
        resource_provenance=iscn_provenance,
        policy_parameters={
            "automatic_unvalidated_sv_to_iscn": False,
            "sex_chromosomes_assessed": False,
            "exact_full_chromosome_span_required_for_iscn": (
                policy.whole_chromosome_span_basis == "exact_contig"
            ),
            "whole_chromosome_span_basis": policy.whole_chromosome_span_basis,
            "cnv_policy_profile_id": policy.profile_id,
            "cnv_policy_schema_version": policy.schema_version,
            "whole_chromosome_fraction": policy.whole_chromosome_fraction,
            "cytoband_affected_fraction": policy.cytoband_affected_fraction,
        },
        technical_assumptions=[
            "Whole-chromosome candidate classification uses the versioned QDNAseq "
            "segment-fraction threshold; +chr/-chr rendering additionally requires a "
            "confirmed whole-chromosome span under the versioned span basis, either an "
            "exact zero-to-contig-end segment or full coverage of that chromosome's "
            "assessable QDNAseq bins.",
            "Segmental fragments require one contiguous affected cytoband group at the "
            "versioned affected-fraction threshold.",
            "Unvalidated SV breakpoint pairs remain review evidence outside formal notation.",
        ],
        assessment_blockers=iscn_blockers,
        cnv_source_event_ids={event.event_id for event in cnv.events},
    )
    iscn_outcome = iscn_module_outcome(proposal)
    return result.model_copy(
        update={
            "events": combined_events,
            "iscn": proposal,
            "modules": _replace_module(_replace_module(result.modules, outcome), iscn_outcome),
            "provenance": _merge_provenance(result.provenance, cnv),
            "warnings": [*result.warnings, *cnv.warnings, *cnv.limitations],
            "sidecars": sidecars,
        }
    )


def report_arguments(ctx: RunContext, cnv: QDNAseqCallReport) -> dict[str, Any]:
    """The keyword arguments with which ``render_html`` draws the copy-number section."""
    return {"cnv_report": cnv, "cnv_evidence_root": ctx.envelope.path(CNV_DIR)}


def enrich_workbook(path: Path, ctx: RunContext, cnv: QDNAseqCallReport) -> None:
    """Add the fit, consensus and primary-segment sheets to the rendered workbook."""
    workbook = load_workbook(path)
    for title in ("CNV Fits", "CNV Consensus", "CNV Segments"):
        if title in workbook.sheetnames:
            del workbook[title]

    fits_sheet = workbook.create_sheet("CNV Fits")
    fits_sheet.append(
        [
            "bin_size_kbp",
            "cellularity",
            "ploidy",
            "fit_error",
            "candidate_count",
            "segment_count",
        ]
    )
    for fit in sorted(cnv.fits, key=lambda value: value.bin_size_kbp):
        fits_sheet.append(
            [
                fit.bin_size_kbp,
                fit.cellularity,
                fit.ploidy,
                fit.fit_error,
                fit.candidate_count,
                fit.segment_count,
            ]
        )

    consensus_sheet = workbook.create_sheet("CNV Consensus")
    consensus_sheet.append(
        [
            "chromosome",
            "median_copy_number",
            "rounded_copy_number",
            "agreeing_bins",
            "contributing_bins",
            "min_copy_number",
            "max_copy_number",
        ]
    )
    for chromosome in cnv.chromosome_consensus:
        consensus_sheet.append(
            [
                chromosome.chromosome,
                chromosome.median_copy_number,
                chromosome.rounded_copy_number,
                chromosome.agreeing_bins,
                chromosome.contributing_bins,
                chromosome.min_copy_number,
                chromosome.max_copy_number,
            ]
        )

    segment_sheet = workbook.create_sheet("CNV Segments")
    segment_path = ctx.envelope.path(f"{CNV_DIR}/{cnv.primary_fit.segment_file}")
    with segment_path.open("r", encoding="utf-8") as handle:
        for line in handle:
            segment_sheet.append(line.rstrip("\n").split("\t"))
    workbook.save(path)
