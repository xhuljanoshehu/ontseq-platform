"""Reviewer presentation of the normalized methylation report (HTML table and workbook).

Presentation adds nothing to the report. What it must get right is what the report already
distinguishes: a region below the coverage floor is not measurable (never 0%), calls that
failed their threshold stay visible, untrusted labels stay text, and a report is only ever
shown next to the result it belongs to.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from ontseq_platform.demo import build_demo_result
from ontseq_platform.methylation import (
    MethylationPolicy,
    MethylationRegionSource,
    MethylationReport,
    _bed_regions,
    normalize_methylation,
)
from ontseq_platform.models import (
    AnalysisModule,
    ModuleOutcome,
    ModuleRunStatus,
    PipelineResult,
    ToolRecord,
)
from ontseq_platform.report import render_html
from ontseq_platform.report_methylation import (
    NOT_MEASURABLE,
    REGION_HEADERS,
    methylation_facts,
    ordered_regions,
    region_rows,
    validate_methylation_identity,
)
from ontseq_platform.workbook import render_workbook

_HOSTILE = "=SUM(1,2)"


def _site(chromosome: str, start: int, code: str, valid: int, modified: int) -> str:
    return (
        f"{chromosome}\t{start}\t{start + 1}\t{code}\t{valid}\t+\t{start}\t{start + 1}\t0,0,0\t"
        f"{valid}\t{0 if valid == 0 else modified * 100 / valid:.2f}\t{modified}\t"
        f"{valid - modified}\t0\t0\t3\t0\t2\n"
    )


def _report(tmp_path: Path) -> MethylationReport:
    bed = tmp_path / "targets.bed"
    bed.write_text(
        "chr10\t0\t100\tTARGET_CHR10\n"
        "chr2\t0\t100\tTARGET_CHR2\n"
        f"chr2\t200\t300\t{_HOSTILE}\n"
        "chr1\t0\t100\t<script>alert(1)</script>\n",
        encoding="utf-8",
    )
    bedmethyl = tmp_path / "sample.bedmethyl"
    bedmethyl.write_text(
        _site("chr10", 10, "m", 20, 15)
        + _site("chr2", 10, "m", 10, 0)  # a measured zero
        + _site("chr2", 11, "h", 10, 1)
        + _site("chr2", 210, "m", 2, 1)  # below the floor of 5: not measurable
        + _site("chr1", 10, "m", 8, 8),
        encoding="utf-8",
    )
    policy = MethylationPolicy(
        profile_id="presentation-test",
        status="technical_defaults_only",
        modification_codes=["m", "h"],
        region_source=MethylationRegionSource.TARGET_BED,
        note="Synthetic presentation fixture, not a validated policy",
    )
    return normalize_methylation(
        sample_id="SYNTHETIC_AML_001",
        genome_build=build_demo_result().manifest.assay.genome_build,
        bedmethyl_path=bedmethyl,
        policy=policy,
        tool=ToolRecord(name="modkit", version="0.6.4"),
        regions=_bed_regions(bed),
        target_bed=bed,
        reads_with_modified_base_tags=None,
    )


def _result_for(report: MethylationReport) -> PipelineResult:
    result = build_demo_result()
    modules = [item for item in result.modules if item.module != AnalysisModule.METHYLATION]
    modules.append(
        ModuleOutcome(
            module=AnalysisModule.METHYLATION,
            status=report.status,
            reason="synthetic presentation fixture",
            tools=[report.tool],
        )
    )
    analysis = result.manifest.analysis.model_copy(
        update={"modules": [*result.manifest.analysis.modules, AnalysisModule.METHYLATION]}
    )
    payload = result.model_copy(
        update={
            "manifest": result.manifest.model_copy(update={"analysis": analysis}),
            "modules": sorted(modules, key=lambda item: item.module.value),
            "provenance": result.provenance.model_copy(
                update={
                    "reference_checksums": {
                        **result.provenance.reference_checksums,
                        "bedmethyl": report.bedmethyl_fingerprint.sha256,
                    }
                }
            ),
        }
    )
    return PipelineResult.model_validate(payload.model_dump(mode="python"))


def test_regions_are_in_genome_order_then_by_modification(tmp_path: Path) -> None:
    ordered = ordered_regions(_report(tmp_path))
    keys = [
        (region.chromosome, region.region_id, region.modification_code.value) for region in ordered
    ]
    assert keys[0][0] == "chr1"
    assert [key[0] for key in keys].index("chr2") < [key[0] for key in keys].index("chr10")
    chr2 = [key for key in keys if key[1] == "TARGET_CHR2"]
    assert [code for *_rest, code in chr2] == ["m", "h"]


def test_a_region_below_the_floor_is_not_measurable_and_never_zero(tmp_path: Path) -> None:
    report = _report(tmp_path)
    rows = {
        (row[0], row[4]): dict(zip(REGION_HEADERS, row, strict=True)) for row in region_rows(report)
    }
    below = rows[(_HOSTILE, "5mC")]
    assert below["Measurement"] == NOT_MEASURABLE
    assert below["Mean modified fraction (call-weighted)"] is None
    assert below["Median site modified fraction"] is None
    assert below["Sites"] == 1 and below["Sites at coverage floor"] == 0
    zero = rows[("TARGET_CHR2", "5mC")]
    assert zero["Measurement"] == "measured"
    assert zero["Mean modified fraction (call-weighted)"] == 0.0
    assert zero["Failed-threshold calls"] == 3 and zero["No-call calls"] == 2


def test_facts_distinguish_an_unanswered_tag_probe_from_zero(tmp_path: Path) -> None:
    facts = dict(methylation_facts(_report(tmp_path)))
    assert facts["Reads with MM tags"].startswith("unknown")
    assert facts["Call confidence threshold (pinned)"] == 0.8
    assert dict(methylation_facts(None))["Status"] == "NOT_RECORDED"


def test_workbook_sheets_keep_empty_fractions_and_text_labels(tmp_path: Path) -> None:
    report = _report(tmp_path)
    path = render_workbook(_result_for(report), tmp_path / "out.xlsx", methylation_report=report)
    workbook = load_workbook(path)
    assert {"13_Methylation", "14_Methylation_Regions"} <= set(workbook.sheetnames)
    sheet = workbook["14_Methylation_Regions"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(REGION_HEADERS)
    fraction_column = headers.index("Mean modified fraction (call-weighted)")
    hostile = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == _HOSTILE)
    assert hostile[0].data_type == "s"
    assert hostile[fraction_column].value is None
    facts = {row[0].value: row[1].value for row in workbook["13_Methylation"].iter_rows(min_row=2)}
    assert facts["Status"] == "COMPLETED"
    assert facts["bedMethyl SHA-256"] == report.bedmethyl_fingerprint.sha256


def test_workbook_without_methylation_has_no_methylation_sheets(tmp_path: Path) -> None:
    path = render_workbook(build_demo_result(), tmp_path / "plain.xlsx")
    assert not any("Methylation" in name for name in load_workbook(path).sheetnames)


def test_requested_methylation_without_a_report_says_so(tmp_path: Path) -> None:
    report = _report(tmp_path)
    result = _result_for(report)
    path = render_workbook(result, tmp_path / "missing.xlsx")
    names = load_workbook(path).sheetnames
    assert "13_Methylation" in names and "14_Methylation_Regions" not in names


def test_html_table_escapes_labels_and_names_unmeasurable_regions(tmp_path: Path) -> None:
    report = _report(tmp_path)
    html = render_html(_result_for(report), tmp_path / "r.html", methylation_report=report)
    document = Path(html).read_text(encoding="utf-8")
    section = document[document.index("<section id='methylation'>") :]
    section = section[: section.index("</section>")]
    assert "Region table" in section
    assert NOT_MEASURABLE in section
    assert "<script>alert(1)</script>" not in section
    assert "&lt;script&gt;" in section
    assert "presentation-test" in section


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"sample_id": "OTHER_SAMPLE"}, "sample/build"),
        ({"bedmethyl_fingerprint": None}, "pileup the result recorded"),
        ({"status": ModuleRunStatus.NO_CALL}, "module outcome"),
    ],
)
def test_a_report_from_elsewhere_is_never_presented(
    tmp_path: Path, change: dict[str, object], message: str
) -> None:
    report = _report(tmp_path)
    result = _result_for(report)
    if "bedmethyl_fingerprint" in change:
        change = {
            "bedmethyl_fingerprint": report.bedmethyl_fingerprint.model_copy(
                update={"sha256": "f" * 64}
            )
        }
    foreign = report.model_copy(update=change)
    with pytest.raises(ValueError, match=message):
        validate_methylation_identity(result, foreign)
    with pytest.raises(ValueError, match=message):
        render_workbook(result, tmp_path / "x.xlsx", methylation_report=foreign)
