from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from ontseq_platform.demo import build_demo_result
from ontseq_platform.report import render_html
from ontseq_platform.report_interactive import interactive_payload
from ontseq_platform.workbook import render_workbook


def test_absent_marlin_remains_unrecorded_in_every_report(tmp_path: Path) -> None:
    result = build_demo_result()
    data = interactive_payload(result, marlin_report=None)
    assert data["marlin"] is None
    document = render_html(result, tmp_path / "report.html", marlin_report=None).read_text()
    assert "Für diesen Lauf kein MARLIN-Ausführungsnachweis vorhanden" in document
    book = load_workbook(render_workbook(result, tmp_path / "report.xlsx", marlin_report=None))
    rows = dict(book["11_MARLIN"].iter_rows(min_row=2, values_only=True))
    assert rows["Status"] == "NOT_RECORDED"
    assert rows["Leading model score"] is None


def native_report(result, status="COMPLETED", score=0.6, label="SYNTHETIC_CLASS"):
    from ontseq_platform.marlin_native_contracts import NativeMarlinReport

    data = dict(
        run_id=result.manifest.run_id,
        sample_id=result.manifest.sample_id,
        genome_build=result.manifest.assay.genome_build,
        status=status,
        reason="Synthetic evidence",
    )
    if status in {"COMPLETED", "NO_CALL"}:
        observed = 10 if status == "COMPLETED" else 0
        data.update(
            feature_summary=dict(
                observed_model_feature_count=observed,
                explicit_na_feature_count=0,
                absent_feature_count=357340 - observed,
                non_model_probe_count=0,
                observed_fraction=observed / 357340,
                feature_vector_sha256="a" * 64,
                feature_artifact_sha256="b" * 64,
            ),
            tools=[dict(name="native synthetic", version="1", parameters={})],
            installation_signature={"model": "c" * 64},
            input_fingerprints={
                key: dict(size_bytes=1, sha256="d" * 64)
                for key in (
                    "bam",
                    "bam_index",
                    "reference_fai",
                    "reference",
                    "bedmethyl",
                    "tensor",
                    "worker_output",
                )
            },
            decision="UNKNOWN",
        )
    if status == "COMPLETED":
        data.update(
            decision="UNKNOWN",
            model_score_threshold_met=score >= 0.8,
            raw_model_scores=[
                dict(model_id=i + 1, label=f"model {i}", score=1 / 42) for i in range(42)
            ],
            class_scores=[dict(label=label, score=score), dict(label="z_other", score=1 - score)],
            family_scores=[dict(label="family", score=1)],
            lineage_scores=[dict(label="lineage", score=1)],
            top_class=label,
            top_class_score=score,
        )
    import hashlib

    from ontseq_platform.models import ModuleOutcome

    report = NativeMarlinReport(**data)
    result.modules = [item for item in result.modules if item.module.value != "marlin"]
    result.modules.append(
        ModuleOutcome(
            module="marlin", status=report.status, reason=report.reason, tools=report.tools
        )
    )
    if status in {"COMPLETED", "NO_CALL"}:
        result.provenance.reference_checksums["marlin_report"] = hashlib.sha256(
            report.model_dump_json().encode()
        ).hexdigest()
    else:
        result.provenance.reference_checksums.pop("marlin_report", None)
    return report


@pytest.mark.parametrize("status", ["NOT_RUN", "FAILED", "NO_CALL", "COMPLETED"])
def test_marlin_status_and_missing_values_survive_exports(tmp_path, status):
    result = build_demo_result()
    report = native_report(result, status)
    data = interactive_payload(result, marlin_report=report)["marlin"]
    assert data["status"] == status
    assert data["top_class_score"] == (0.6 if status == "COMPLETED" else None)
    doc = render_html(result, tmp_path / "report.html", marlin_report=report).read_text()
    assert status in doc and "UNVALIDATED_RESEARCH" in doc
    assert "SECRET" not in doc
    book = load_workbook(render_workbook(result, tmp_path / "report.xlsx", marlin_report=report))
    rows = dict(book["11_MARLIN"].iter_rows(min_row=2, values_only=True))
    assert rows["Status"] == status
    assert rows["Observed model CpGs"] == (
        10 if status == "COMPLETED" else 0 if status == "NO_CALL" else None
    )
    assert rows["Leading model score"] == (0.6 if status == "COMPLETED" else None)


@pytest.mark.parametrize(
    "key,value", [("sample_id", "other"), ("run_id", "other"), ("genome_build", "GRCh37")]
)
def test_marlin_rejects_cross_sample_build_run(tmp_path, key, value):
    result = build_demo_result()
    report = native_report(result, "NOT_RUN")
    if key == "genome_build" and str(report.genome_build) == value:
        value = "GRCh38"
    report = report.model_copy(update={key: value})
    for render in (render_html, render_workbook):
        with pytest.raises(ValueError, match="identity"):
            render(result, tmp_path / "bad", marlin_report=report)
    with pytest.raises(ValueError, match="identity"):
        interactive_payload(result, marlin_report=report)


def test_unknown_keeps_leading_score_without_confirmed_class_and_escapes_labels(tmp_path):
    result = build_demo_result()
    report = native_report(result, label="<script>unsafe()</script>")
    doc = render_html(result, tmp_path / "report.html", marlin_report=report).read_text()
    assert "UNKNOWN · Keine hinreichend sichere Klassifikation" in doc
    assert "Leading model score" in doc and "0.6" in doc
    assert "<script>unsafe()</script>" not in doc
    assert "&lt;script&gt;unsafe()&lt;/script&gt;" in doc
    assert "keine Erkrankungswahrscheinlichkeiten" in doc
    report = native_report(result, score=0.9, label="=UNTRUSTED()")
    book = load_workbook(render_workbook(result, tmp_path / "report.xlsx", marlin_report=report))
    assert book["12_MARLIN_Scores"]["C2"].value == "=UNTRUSTED()"
    assert book["12_MARLIN_Scores"]["C2"].data_type == "s"
    rows = dict(book["11_MARLIN"].iter_rows(min_row=2, values_only=True))
    assert rows["Decision"] == "UNKNOWN"
    assert rows["Assay assessability"] == "NOT_ESTABLISHED"


def archived_run(root, status="COMPLETED"):
    import hashlib
    from datetime import UTC, datetime

    from ontseq_platform.pipeline.runner import RUN_REPORT
    from ontseq_platform.pipeline.state import ArtifactRecord, RunReport, StageRecord

    result = build_demo_result()
    report = native_report(result, status)
    records = []
    if status in {"COMPLETED", "NO_CALL"}:
        relative = f"evidence/marlin/{result.manifest.sample_id}.marlin.json"
        path = root / relative
        path.parent.mkdir(parents=True)
        path.write_text(report.model_dump_json())
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        result.provenance.reference_checksums["marlin_report"] = digest
        records.append(
            ArtifactRecord(
                relative_path=relative,
                size_bytes=path.stat().st_size,
                sha256=digest,
                exportable=True,
            )
        )
    path = root / f"normalized/{result.manifest.sample_id}.result.json"
    path.parent.mkdir(parents=True)
    path.write_text(result.model_dump_json())
    normalized = ArtifactRecord(
        relative_path=path.relative_to(root).as_posix(),
        size_bytes=path.stat().st_size,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        exportable=True,
    )
    stages = [
        StageRecord(
            stage="marlin",
            title="MARLIN",
            status=status,
            verification="verified_pure_python",
            required=False,
            reason=report.reason,
            signature="0" * 64 if records else None,
            outputs=records,
            tools=report.tools,
        ),
        StageRecord(
            stage="assemble",
            title="assemble",
            status="COMPLETED",
            verification="verified_pure_python",
            required=True,
            reason="synthetic",
            signature="0" * 64,
            outputs=[normalized],
        ),
    ]
    run = RunReport(
        run_id=result.manifest.run_id,
        sample_id=result.manifest.sample_id,
        input_kind=result.manifest.input.kind,
        genome_build=result.manifest.assay.genome_build,
        manifest=result.manifest,
        passed=True,
        verdict_reason="synthetic",
        pipeline_version="0.8.2",
        git_commit="synthetic",
        started_at=datetime.now(UTC),
        finished_at=datetime.now(UTC),
        stages=stages,
    )
    runpath = root / RUN_REPORT
    runpath.parent.mkdir(parents=True, exist_ok=True)
    runpath.write_text(run.model_dump_json())
    return result, report, run


@pytest.mark.parametrize("status", ["COMPLETED", "NO_CALL", "FAILED", "NOT_RUN"])
def test_archive_export_binds_marlin_status_to_run(tmp_path, status):
    from test_interactive_report import payload

    from ontseq_platform.service.befund import current_befund

    result, report, run = archived_run(tmp_path, status)
    assert payload(current_befund(tmp_path, result).decode())["marlin"]["status"] == status


@pytest.mark.parametrize("tamper", ["checksum", "outcome", "unrecorded", "reference_hash"])
def test_archive_export_rejects_stale_marlin(tmp_path, tamper):
    from ontseq_platform.pipeline.runner import RUN_REPORT
    from ontseq_platform.service.befund import current_befund

    result, report, run = archived_run(tmp_path)
    if tamper == "checksum":
        (tmp_path / run.stages[0].outputs[0].relative_path).write_text(
            report.model_dump_json() + " "
        )
    elif tamper == "outcome":
        result.modules[-1].reason = "stale normalized result"
    elif tamper == "unrecorded":
        run.stages.pop(0)
        (tmp_path / RUN_REPORT).write_text(run.model_dump_json())
    else:
        result.provenance.reference_checksums["marlin_report"] = "f" * 64
    with pytest.raises(ValueError, match="MARLIN|checksum|evidence"):
        current_befund(tmp_path, result)


def test_archive_accepts_resumed_stage_reason_prefix(tmp_path):
    from test_interactive_report import payload

    from ontseq_platform.pipeline.runner import RUN_REPORT
    from ontseq_platform.service.befund import current_befund

    result, report, run = archived_run(tmp_path)
    run.stages[0].reason = "Resumed unchanged from a previous run. " * 2 + report.reason
    run.stages[0].resumed = True
    (tmp_path / RUN_REPORT).write_text(run.model_dump_json())
    assert (
        payload(current_befund(tmp_path, result).decode())["marlin"]["top_class"]
        == report.top_class
    )


def test_archive_unrequested_stage_is_absent_not_an_export_error(tmp_path):
    from test_interactive_report import payload

    from ontseq_platform.pipeline.runner import RUN_REPORT
    from ontseq_platform.service.befund import current_befund

    result, report, run = archived_run(tmp_path, "NOT_RUN")
    result.modules = [item for item in result.modules if item.module.value != "marlin"]
    result.manifest.analysis.modules = [
        item
        for item in result.manifest.analysis.modules
        if item.value not in {"marlin", "methylation"}
    ]
    run.manifest = result.manifest
    import hashlib

    result_path = tmp_path / f"normalized/{result.manifest.sample_id}.result.json"
    result_path.write_text(result.model_dump_json())
    record = run.stages[1].outputs[0]
    record.size_bytes = result_path.stat().st_size
    record.sha256 = hashlib.sha256(result_path.read_bytes()).hexdigest()
    (tmp_path / RUN_REPORT).write_text(run.model_dump_json())
    assert payload(current_befund(tmp_path, result).decode())["marlin"] is None


@pytest.mark.parametrize("status", ["NOT_RUN", "FAILED"])
def test_archive_failed_or_blocked_ignores_unrecorded_stale_file(tmp_path, status):
    from test_interactive_report import payload

    from ontseq_platform.service.befund import current_befund

    result, report, run = archived_run(tmp_path, status)
    path = tmp_path / f"evidence/marlin/{result.manifest.sample_id}.marlin.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("STALE INVALID CONTENT MUST NEVER BE PARSED")
    data = payload(current_befund(tmp_path, result).decode())["marlin"]
    assert data["status"] == status and data["top_class"] is None


@pytest.mark.parametrize(
    "contradiction", ["FAILED", "NOT_RUN", "absent", "tools", "reason", "missing_hash"]
)
def test_public_report_rejects_conflicting_marlin_outcome(tmp_path, contradiction):
    from ontseq_platform.models import ModuleOutcome, ModuleRunStatus

    result = build_demo_result()
    report = native_report(result)
    result.modules = [item for item in result.modules if item.module.value != "marlin"]
    result.modules.append(
        ModuleOutcome(
            module="marlin", status=report.status, reason=report.reason, tools=report.tools
        )
    )
    result.provenance.reference_checksums["marlin_report"] = "a" * 64
    if contradiction in {"FAILED", "NOT_RUN"}:
        result.modules[-1].status = ModuleRunStatus(contradiction)
    elif contradiction == "absent":
        result.modules.pop()
    elif contradiction == "tools":
        result.modules[-1].tools = []
    elif contradiction == "reason":
        result.modules[-1].reason = "another reason"
    else:
        result.provenance.reference_checksums.pop("marlin_report")
    for render in (render_html, render_workbook):
        with pytest.raises(ValueError, match="MARLIN"):
            render(result, tmp_path / "bad", marlin_report=report)
    with pytest.raises(ValueError, match="MARLIN"):
        interactive_payload(result, marlin_report=report)
