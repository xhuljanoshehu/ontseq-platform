"""Native MARLIN score thresholds cannot establish assay assessability."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_marlin_report_presentation import native_report

from ontseq_platform.demo import build_demo_result
from ontseq_platform.report import render_html
from ontseq_platform.report_interactive import interactive_payload
from ontseq_platform.report_marlin import marlin_facts
from ontseq_platform.workbook import render_workbook


def test_report_separates_observed_fraction_threshold_and_assay_assessability(tmp_path: Path):
    result = build_demo_result()
    report = native_report(result)
    facts = dict(marlin_facts(report))
    assert facts["Observed model CpG fraction"] == 10 / 357340
    assert facts["Raw model score threshold"] == 0.8
    assert facts["Model score threshold met"] is False
    assert facts["Assay assessability"] == "NOT_ESTABLISHED"
    assert facts["Assessment"] == "Bewertung nicht validiert"
    html = render_html(result, tmp_path / "report.html", marlin_report=report).read_text()
    assert "Bewertung nicht validiert" in html
    book = load_workbook(render_workbook(result, tmp_path / "report.xlsx", marlin_report=report))
    rows = dict(book["11_MARLIN"].iter_rows(min_row=2, values_only=True))
    assert rows["Assay assessability"] == "NOT_ESTABLISHED"
    assert rows["Model score threshold met"] is False
    data = interactive_payload(result, marlin_report=report)["marlin"]
    assert data["model_score_threshold_met"] is False
    assert data["assay_assessability"] == "NOT_ESTABLISHED"


@pytest.mark.parametrize("observed", [1, 357340])
def test_high_model_score_remains_unknown_at_any_feature_coverage(tmp_path, observed):
    import hashlib

    from ontseq_platform.marlin_native_contracts import NativeMarlinReport

    result = build_demo_result()
    original = native_report(result, score=0.9)
    data = original.model_dump(mode="json")
    data["feature_summary"].update(
        observed_model_feature_count=observed,
        absent_feature_count=357340 - observed,
        observed_fraction=observed / 357340,
    )
    report = NativeMarlinReport.model_validate(data)
    result.provenance.reference_checksums["marlin_report"] = hashlib.sha256(
        report.model_dump_json().encode()
    ).hexdigest()
    html = render_html(result, tmp_path / "report.html", marlin_report=report).read_text()
    assert "HIGH_CONFIDENCE" not in html
    assert "Bewertung nicht validiert" in html
    book = load_workbook(render_workbook(result, tmp_path / "report.xlsx", marlin_report=report))
    facts = dict(book["11_MARLIN"].iter_rows(min_row=2, values_only=True))
    assert facts["Decision"] == "UNKNOWN"
    assert facts["Model score threshold met"] is True
    assert facts["Leading model score"] == 0.9
    assert facts["Observed model CpG fraction"] == pytest.approx(
        observed / 357340, rel=1e-14, abs=0
    )
    assert facts["Assay assessability"] == "NOT_ESTABLISHED"


def test_no_call_does_not_present_an_unmeasured_threshold_as_false():
    result = build_demo_result()
    report = native_report(result, "NO_CALL")
    facts = dict(marlin_facts(report))
    assert facts["Observed model CpG fraction"] == 0
    assert facts["Model score threshold met"] is None
    assert facts["Status"] == "NO_CALL"
